from io import BytesIO
import json
from pathlib import Path
from types import SimpleNamespace
import unittest
from unittest.mock import patch

from openpyxl import load_workbook

from core import context, demo_outputs, llm, orchestrator, parser, plan, prompts, qa, schema, scoring


ROOT = Path(__file__).resolve().parents[1]
STRESS_FILE = ROOT / "data" / "go_japan_campaign_stress_test.xlsx"


class SchemaPipelineTests(unittest.TestCase):
    def _tables_and_flags(self):
        tables, _ = parser.load_workbook_with_report(STRESS_FILE)
        flags = [dict(flag, _id=f"C-{index}") for index, flag in enumerate(
            scoring.dedupe_flags(parser.validate(tables) + scoring.detect_conflicts(tables)), 1
        )]
        return tables, flags

    def _valid_structured_market(self, tables, market):
        contract = next(item for item in plan.recommendation_contract(tables)["markets"]
                        if item["market"] == market)
        channel = contract["channels"]["eligible"][0]["_name"]
        reading = {
            "recommendation": "Prioritize a focused staged activation plan",
            "rationale": "Market evidence shows a specific conversion barrier that this sequencing directly addresses.",
            "evidence_refs": ["B.MAIN_USER_BARRIER"],
        }
        reading_refs = {
            "positioning": ["B.MAIN_USER_BARRIER"],
            "city_portfolio": [f"DERIVED.CITY_ROLE.{orchestrator._slug(contract['cities'][0]['city'])}"],
            "modules": ["B.MAIN_USER_BARRIER"],
            "crm": ["B.BOOKING_LEAD_TIME", f"E.CHANNEL.{orchestrator._slug(channel)}"],
            "assets": ["H.CONSTRAINT.REGIONAL_TEAMS"],
            "localization": ["B.LOCALIZATION_NEED"],
        }
        cities = []
        for city in contract["cities"]:
            resources = city["role"] != "NEEDS DATA"
            cities.append({
                "city": city["city"],
                "product_angle": "Use a focused evidence-led product angle",
                "play": "Sequence education before the conversion action",
                "modules": ([{"id": "Hero KV", "justification": "Establishes the market proposition"}]
                            if resources else []),
                "channels": ([{"id": channel, "justification": "Matches the journey stage"}]
                             if resources else []),
                "assets": [],
            })
        product_ids = orchestrator._product_ids(contract)
        priority_item = lambda item_id, ref: {
            "id": item_id,
            "decision": "Use this item at its recommended journey stage",
            "reason": "Its source evidence supports this relative execution order",
            "evidence_refs": [ref],
        }
        return {
            "market": market,
            "readings": {name: {**reading, "evidence_refs": reading_refs[name]} for name in
                         ("positioning", "city_portfolio", "modules", "crm", "assets", "localization")},
            "market_summary": "Sequence the strongest cities before conditional opportunities",
            "immediate_next_action": "Approve the first-stage market activation plan",
            "product_focus": "Prioritize the strongest complete product signals available",
            "stakeholder_constraints": [{
                "stakeholder": "Regional Teams",
                "impact": "Regional review controls localization before activation",
                "applies_to": ["localization"],
            }],
            "priority_plan": {
                "cities": [{**priority_item(
                    city["city"], f"DERIVED.CITY_ROLE.{orchestrator._slug(city['city'])}"
                ), "decision": ({
                    "Conditional Conversion": "Confirm supply before conditional conversion activation",
                    "Package": "Use a packaged tour approach for this city",
                    "Content": "Keep this city as content inspiration only",
                    "NEEDS DATA": "Hold activation and obtain readiness data first",
                }.get(city["role"], "Use this city at its recommended journey stage"))}
                    for city in contract["cities"]],
                "products": [priority_item(
                    product, f"C.{orchestrator._slug(product)}"
                ) for product in product_ids],
                "channels": [priority_item(
                    channel, f"E.CHANNEL.{orchestrator._slug(channel)}"
                )],
                "modules": [priority_item("Hero KV", "F.MODULE.HERO_KV")],
                "assets": [],
            },
            "crm": {
                "objective": "Move qualified audiences through a staged conversion journey",
                "rationale": "The market barrier requires education before product conversion",
                "evidence_refs": ["B.BOOKING_LEAD_TIME", "B.MAIN_USER_BARRIER"],
                "segments": [{"name": "Recent destination researchers",
                              "window": "illustrative 30-day window",
                              "reason": "Build intent before presenting conversion modules",
                              "assumption": True}],
                "weekly_plan": [
                    {"week": week, "objective": f"Advance campaign stage number {week}",
                     "audience": "Recent destination research audience", "channel_ids": [channel],
                     "action": f"Deliver the planned stage {week} market message", "gate": "",
                     "evidence_refs": ["B.BOOKING_LEAD_TIME",
                                       f"E.CHANNEL.{orchestrator._slug(channel)}"]}
                    for week in orchestrator.campaign_weeks(tables)
                ],
                "push_copy": [], "edm_copy": [],
            },
            "cities": cities,
        }

    def test_destination_header_aliases_are_canonicalized(self):
        tables, report = parser.load_workbook_with_report(STRESS_FILE)

        self.assertFalse(report["blocking"])
        self.assertEqual(tables["B_Markets"][0]["Search UV Index"], "92")
        self.assertNotIn("Japan Search UV Index", tables["B_Markets"][0])
        self.assertEqual(set(tables["F_Modules"][0]), {
            "Module", "Description", "CTR Index", "Conversion Contribution",
            "Design Complexity", "Notes",
        })

    def test_unknown_header_at_fixed_position_blocks_analysis(self):
        workbook = load_workbook(STRESS_FILE)
        workbook["B_Market_Demand"].cell(1, 2).value = "Demand Score"
        payload = BytesIO()
        workbook.save(payload)
        payload.seek(0)

        _, report = parser.load_workbook_with_report(payload)

        self.assertTrue(report["blocking"])
        self.assertTrue(any(
            issue["code"] == "HEADER_MISMATCH" and issue["table"] == "B_Markets"
            for issue in report["issues"]
        ))

    def test_missing_numeric_input_is_not_coerced_to_zero(self):
        tables, _ = parser.load_workbook_with_report(STRESS_FILE)
        scores = {row["market"]: row for row in scoring.market_scores(tables)}

        self.assertIsNone(scores["India"]["score"])
        self.assertEqual(scores["India"]["tier"], "NEEDS DATA")

    def test_scoring_table_hides_internal_notes_column(self):
        tables, _ = parser.load_workbook_with_report(ROOT / "data" / "go_china.xlsx")
        rendered = scoring.scores_md(scoring.market_scores(tables), collapse_methodology=True)

        self.assertNotIn("| Notes |", rendered)
        self.assertIn("| Market | UV_norm | Growth_norm | Product_norm (mix) | Score | Tier |", rendered)

    def test_demo_contains_current_global_brief_and_all_market_packs(self):
        expected = {"__global__", "Hong Kong", "Korea", "Singapore", "Malaysia", "Thailand"}

        self.assertEqual(set(demo_outputs.DEMO), expected)
        self.assertTrue(demo_outputs.DEMO["__global__"].startswith("# Go China Round 7"))
        for market in expected - {"__global__"}:
            pack = demo_outputs.DEMO[market]
            self.assertTrue(pack.startswith(f"# MARKET EXECUTION PACK: {market}"))
            self.assertNotIn("| Notes |", pack)
            self.assertNotIn("| Product | — |", pack)
            asset_section = pack.split("## 5. Asset requests", 1)[1].split(
                "## 6. Localization checklist", 1
            )[0]
            self.assertLess(asset_section.index("Applicable stakeholder constraints"),
                            asset_section.index("| Asset | Priority | Why selected [AI REC] |"))

    def test_demo_outputs_are_brand_neutral_and_free_of_corrupted_separators(self):
        forbidden = ("trip" + ".com", "c" + "trip", "\u643a\u7a0b", "trip-" + "logo")

        for name, document in demo_outputs.DEMO.items():
            normalized = document.lower()
            self.assertNotIn("\u0431\u043a", document, name)
            for marker in forbidden:
                self.assertNotIn(marker, normalized, name)

    def test_market_codes_and_readiness_gaps_are_detected(self):
        tables, _ = parser.load_workbook_with_report(STRESS_FILE)
        flags = parser.validate(tables) + scoring.detect_conflicts(tables)

        self.assertFalse(any(flag["type"] == "UNMAPPED_MARKET" for flag in flags))
        gaps = {(flag.get("market"), flag.get("city")) for flag in flags
                if flag["type"] == "READINESS_GAP"}
        self.assertEqual(gaps, {("Australia", "Okinwa"), ("India", "Matsuyama")})

    def test_matrix_contract_locks_roles_and_asset_priorities(self):
        tables, _ = parser.load_workbook_with_report(STRESS_FILE)
        contract = plan.recommendation_contract_md(tables)

        self.assertIn("Australia (code AU", contract)
        self.assertIn("Okinwa = NEEDS DATA", contract)
        self.assertIn("City Module Image [P1]", contract)
        self.assertIn("CRM Header [P1]", contract)
        self.assertIn("Content Cover Image [P2]", contract)
        self.assertIn("Creator Short Video [P-1]", contract)
        self.assertIn("AMBIGUOUS - contribution missing/TBD", contract)

    def test_qa_catches_core_generation_failures(self):
        tables, _ = parser.load_workbook_with_report(STRESS_FILE)
        bad = """## 9. Campaign recommendation matrix
| Market | City | Locked role | Channels | Creative assets |
|---|---|---|---|---|
| Australia | Okinwa | Conversion | App Push | City Module Image (P0) |

## 4. CRM plan
Campaign Week 1 and Campaign Week 2. Push copy example 1: Limited seats.
"""
        warnings = qa.lint_all(bad, 5, tables, market="Australia")

        self.assertTrue(any("Incomplete CRM calendar" in warning for warning in warnings))
        self.assertTrue(any("Asset priority mismatch" in warning for warning in warnings))
        self.assertTrue(any("Ambiguous-channel copy" in warning for warning in warnings))

    def test_source_backed_promotion_language_is_not_overconstrained(self):
        tables = {"A_Campaign": {"Campaign Duration": "2 weeks",
                                  "Offer note": "Limited seats"}}
        warnings = qa.lint("Limited seats [FACT - table A]", 2, tables)

        self.assertFalse(any("limited seats" in warning.lower() for warning in warnings))

    def test_market_context_includes_global_campaign_configuration(self):
        tables, _ = parser.load_workbook_with_report(STRESS_FILE)
        market_context, trace = context.assemble_market(
            tables, "Australia", [], upstream=""
        )

        self.assertIn("Campaign configuration (table A, global)", market_context)
        self.assertIn(str(tables["A_Campaign"]["Campaign Duration"]), market_context)
        self.assertIn(str(tables["A_Campaign"]["Campaign Objective"]), market_context)
        self.assertEqual(
            trace["A campaign configuration"],
            "full (global campaign scope, objective and duration)",
        )

    def test_section_three_is_a_market_level_summary(self):
        tables, flags = self._tables_and_flags()
        markets = parser.parse_markets(tables)
        recommendations = {
            market: orchestrator.normalize_market_recommendation(
                {"cities": []}, tables, market, flags
            )[0]
            for market in markets
        }
        synthesis = orchestrator.normalize_global_synthesis({}, markets)
        brief = orchestrator.render_global_brief(tables, recommendations, synthesis, flags)
        section = brief.split("## 3. Executive market x city summary", 1)[1].split(
            "## 4. Product focus by market", 1
        )[0]

        self.assertIn("Immediate next action [AI REC]", section)
        for market in markets:
            self.assertEqual(section.count(f"| {market} |"), 1)

    def test_market_pack_exposes_the_required_abstraction(self):
        tables, flags = self._tables_and_flags()
        raw = self._valid_structured_market(tables, "Australia")
        recommendation, _ = orchestrator.normalize_market_recommendation(
            raw, tables, "Australia", flags
        )
        pack = orchestrator.render_market_pack(tables, recommendation, flags)
        summary = pack.split("## 1. Market positioning & decision-chain summary", 1)[1].split(
            "### Execution priority", 1
        )[0]

        self.assertIn("City / readiness cluster [CODE-LOCKED]", summary)
        self.assertIn("Product focus [AI REC]", summary)
        self.assertIn("Channel choice [AI REC]", summary)
        self.assertIn("Asset choice [AI REC]", summary)
        self.assertIn("Stakeholder constraint", summary)
        self.assertEqual(sum(line.startswith("| Australia |")
                             for line in summary.splitlines()), 1)

    def test_structured_normalization_restores_locked_rows_and_filters_resources(self):
        tables, flags = self._tables_and_flags()
        raw = {
            "market": "Australia renamed by model",
            "positioning": "Seasonal value-led Japan escape",
            "cities": [
                {
                    "city": "Tokyo",
                    "cluster": "Model invented cluster",
                    "product_angle": "Flight and hotel bundle",
                    "play": "Lead with seasonal value",
                    "modules": ["Unknown Module"],
                    "channels": ["EDM"],
                    "assets": ["Content Cover Image", "Main KV"],
                },
                {"city": "Osaka", "product_angle": "Food", "play": "City break"},
                {"city": "Sapporo", "product_angle": "Snow", "play": "Seasonal"},
            ],
        }

        recommendation, warnings = orchestrator.normalize_market_recommendation(
            raw, tables, "Australia", flags
        )

        self.assertEqual(recommendation["market"], "Australia")
        self.assertEqual(
            [city["city"] for city in recommendation["cities"]],
            ["Tokyo", "Osaka", "Sapporo", "Okinwa"],
        )
        by_city = {city["city"]: city for city in recommendation["cities"]}
        self.assertEqual(by_city["Osaka"]["cluster"], "Osaka / Kyoto")
        self.assertEqual(by_city["Sapporo"]["cluster"], "Hokkaido / Sapporo")
        self.assertEqual(by_city["Okinwa"]["role"], "NEEDS DATA")
        self.assertEqual(by_city["Okinwa"]["channels"], [])
        self.assertEqual(by_city["Tokyo"]["channels"], [])
        self.assertEqual(
            by_city["Tokyo"]["assets"],
            [{"id": "Main KV", "priority": "P0", "conditional": False,
              "justification": ""}],
        )
        self.assertTrue(any("Okinwa: model row missing" in warning for warning in warnings))
        self.assertTrue(any("ineligible channel 'EDM'" in warning for warning in warnings))
        self.assertTrue(any("P2 asset 'Content Cover Image'" in warning for warning in warnings))

    def test_code_assembled_matrix_contains_every_interest_city(self):
        tables, flags = self._tables_and_flags()
        markets = parser.parse_markets(tables)
        recommendations = {
            market: orchestrator.normalize_market_recommendation(
                {"cities": []}, tables, market, flags
            )[0]
            for market in markets
        }
        synthesis = orchestrator.normalize_global_synthesis({}, markets)

        brief = orchestrator.render_global_brief(
            tables, recommendations, synthesis, flags
        )

        australia_rows = [line for line in brief.splitlines()
                          if line.startswith("| Australia |")]
        self.assertTrue(any("| Okinwa |" not in line and "Okinwa" in line
                            for line in australia_rows))
        self.assertIn("Sapporo (readiness cluster: Hokkaido / Sapporo)", brief)
        self.assertIn("Okinwa | NEEDS DATA", brief)
        self.assertIn("<summary>Scoring methodology</summary>", brief)
        self.assertIn("| Market | Product focus [AI REC] | Applicable stakeholder constraints |", brief)
        self.assertNotIn("Evidence status", brief)
        self.assertNotIn("Complete scoring inputs", brief)
        self.assertIn("| Checklist item | Owner | Why it needs confirmation | Blocking ID |", brief)
        self.assertNotIn("Brief drafted", brief)
        self.assertNotIn("Pending validation", brief)
        self.assertEqual(
            qa.lint_required_ids(brief, [flag["_id"] for flag in flags]), []
        )
        warnings = qa.lint_all(brief, 5, tables)
        self.assertFalse(any("Matrix coverage missing" in warning for warning in warnings))
        self.assertFalse(any("City role mismatch" in warning for warning in warnings))

    def test_market_pack_reuses_validated_recommendation_and_local_ids(self):
        tables, flags = self._tables_and_flags()
        recommendation, _ = orchestrator.normalize_market_recommendation(
            {"cities": []}, tables, "Australia", flags
        )

        pack = orchestrator.render_market_pack(tables, recommendation, flags)

        self.assertIn("City / readiness cluster [CODE-LOCKED]", pack)
        self.assertIn("| AU-1 |", pack)
        self.assertIn("Blocked by C-", pack)
        self.assertNotIn("| C-1 |", pack)

    def test_json_extraction_accepts_fences_and_rejects_arrays(self):
        self.assertEqual(orchestrator.extract_json('```json\n{"ok": true}\n```'), {"ok": True})
        with self.assertRaises(orchestrator.StructuredOutputError):
            orchestrator.extract_json("[]")

    def test_confirmed_coverage_alias_drives_validation_and_market_slicing(self):
        tables, _ = parser.load_workbook_with_report(STRESS_FILE)
        channel = next(row for row in tables["E_Channels"]
                       if row["Channel"] == "Content Channel")
        asset = next(row for row in tables["G_Assets"]
                     if row["Asset Type"] == "Content Cover Image")
        channel["Market Coverage"] = channel["Market Coverage"].replace("AU", "AUS")
        asset["Required Markets"] = asset["Required Markets"].replace("AU", "AUS")
        schema.clear_session_aliases()
        try:
            self.assertEqual(schema.suggest_market_alias("AUS", parser.parse_markets(tables)),
                             "Australia")
            self.assertTrue(any(flag["type"] == "UNMAPPED_MARKET"
                                for flag in parser.validate(tables)))

            schema.register_alias("Australia", "AUS")
            self.assertFalse(any(flag["type"] == "UNMAPPED_MARKET"
                                 for flag in parser.validate(tables)))
            _, trace = context.assemble_market(tables, "Australia", [], upstream="")
            self.assertIn("Content Channel", trace["E rows (coverage filter)"])
        finally:
            schema.clear_session_aliases()

    def test_semantic_validator_rejects_incomplete_crm_instead_of_padding_weeks(self):
        tables, _ = self._tables_and_flags()
        raw = self._valid_structured_market(tables, "Australia")
        raw["crm"]["weekly_plan"] = raw["crm"]["weekly_plan"][:2]

        issues = orchestrator.validate_market_json(raw, tables, "Australia")

        self.assertTrue(any("exactly weeks" in issue for issue in issues))

    def test_semantic_validator_rejects_fabricated_evidence_ids(self):
        tables, _ = self._tables_and_flags()
        raw = self._valid_structured_market(tables, "Australia")
        raw["readings"]["crm"]["evidence_refs"] = ["MODEL.INVENTED.EVIDENCE"]

        issues = orchestrator.validate_market_json(raw, tables, "Australia")

        self.assertTrue(any("invalid IDs" in issue for issue in issues))

    def test_semantic_validator_allows_all_relevant_evidence_without_numeric_cap(self):
        tables, _ = self._tables_and_flags()
        raw = self._valid_structured_market(tables, "Australia")
        channel = raw["crm"]["weekly_plan"][0]["channel_ids"][0]
        raw["crm"]["weekly_plan"][0]["evidence_refs"] = [
            "A.CAMPAIGN_DURATION",
            "A.CAMPAIGN_OBJECTIVE",
            "B.BOOKING_LEAD_TIME",
            "B.MAIN_USER_BARRIER",
            f"E.CHANNEL.{orchestrator._slug(channel)}",
        ]

        issues = orchestrator.validate_market_json(raw, tables, "Australia")

        self.assertFalse(any("weekly_plan[1].evidence_refs" in issue for issue in issues))

    def test_lint_does_not_confuse_clusters_or_channel_names_with_assets(self):
        tables, _ = parser.load_workbook_with_report(ROOT / "data" / "go_china.xlsx")
        text = """| Korea | Confirm Chengdu / Chongqing product depth |\n
| Korea | Shanghai | Conversion | App Homepage Banner | City Module Image [P1] |"""

        warnings = qa.lint_market_city(text, tables) + qa.lint_asset_priorities(text, tables)

        self.assertFalse(any("Chongqing" in warning for warning in warnings))
        self.assertFalse(any("Homepage Banner" in warning for warning in warnings))

    def test_targeted_repair_produces_complete_crm_and_decision_readings(self):
        tables, flags = self._tables_and_flags()
        invalid = self._valid_structured_market(tables, "Australia")
        invalid["crm"]["weekly_plan"] = invalid["crm"]["weekly_plan"][:2]
        valid = self._valid_structured_market(tables, "Australia")
        responses = iter((json.dumps(invalid), json.dumps(valid)))
        calls = []

        raw = orchestrator.call_validated_json(
            lambda prompt: calls.append(prompt) or next(responses),
            "original market prompt",
            lambda value: orchestrator.validate_market_json(value, tables, "Australia"),
        )
        recommendation, _ = orchestrator.normalize_market_recommendation(
            raw, tables, "Australia", flags
        )
        pack = orchestrator.render_market_pack(tables, recommendation, flags)

        self.assertEqual(len(calls), 2)
        self.assertIn("failed validation", calls[1])
        self.assertNotIn("PM to calibrate", pack)
        self.assertIn("| Week | Objective | Audience | Channel | Action | Gate | Decision basis |", pack)
        self.assertNotIn("Therefore,", pack)
        self.assertNotIn("**Owner:**", pack)
        self.assertIn("<summary>Decision basis</summary>", pack)
        self.assertIn("Booking Lead Time =", pack)
        self.assertNotIn("B.BOOKING_LEAD_TIME", pack)
        self.assertIn("| Checklist item | Owner | Why it needs confirmation | Blocking ID |", pack)
        self.assertNotIn("Brief drafted", pack)
        self.assertNotIn("Plan drafted", pack)
        module_checklist = next(
            line for line in pack.splitlines()
            if line.startswith("| Implement selected page modules |")
        )
        self.assertEqual(module_checklist.count("Hero KV"), 1)

    def test_priority_plan_is_validated_and_model_order_is_preserved(self):
        tables, flags = self._tables_and_flags()
        raw = self._valid_structured_market(tables, "Australia")
        raw["priority_plan"]["cities"] = list(reversed(raw["priority_plan"]["cities"]))

        self.assertEqual(orchestrator.validate_market_json(raw, tables, "Australia"), [])
        recommendation, _ = orchestrator.normalize_market_recommendation(
            raw, tables, "Australia", flags
        )
        pack = orchestrator.render_market_pack(tables, recommendation, flags)

        self.assertEqual(recommendation["priority_plan"]["cities"][0]["id"], "Okinwa")
        self.assertIn("### Execution priority [AI REC]", pack)
        self.assertLess(pack.index("| City | 1 | Okinwa |"),
                        pack.index("| City | 4 | Tokyo |"))
        self.assertLess(pack.index("1. **Okinwa**"), pack.index("4. **Tokyo**"))

        raw["priority_plan"]["channels"].append({
            "id": "EDM", "decision": "Use an extra unselected channel in activation",
            "reason": "This deliberately violates the selected resource contract",
            "evidence_refs": ["E.CHANNEL.EDM"],
        })
        issues = orchestrator.validate_market_json(raw, tables, "Australia")
        self.assertFalse(any("priority_plan" in issue for issue in issues))
        normalized, _ = orchestrator.normalize_market_recommendation(
            raw, tables, "Australia", flags
        )
        self.assertNotIn("EDM", [item["id"] for item in normalized["priority_plan"]["channels"]])

    def test_missing_product_decisions_trigger_existing_semantic_repair(self):
        tables, flags = self._tables_and_flags()
        invalid = self._valid_structured_market(tables, "Australia")
        invalid["priority_plan"]["products"] = []
        valid = self._valid_structured_market(tables, "Australia")
        responses = iter((json.dumps(invalid), json.dumps(valid)))
        calls = []

        raw = orchestrator.call_validated_json(
            lambda prompt: calls.append(prompt) or next(responses),
            "original market prompt",
            lambda value: orchestrator.validate_market_json(value, tables, "Australia"),
        )
        recommendation, _ = orchestrator.normalize_market_recommendation(
            raw, tables, "Australia", flags
        )
        pack = orchestrator.render_market_pack(tables, recommendation, flags)

        self.assertEqual(len(calls), 2)
        self.assertIn("priority_plan.products", calls[1])
        self.assertNotIn("AI priority decision unavailable", pack)
        self.assertTrue(all(
            item["model_recommended"]
            for item in recommendation["priority_plan"]["products"]
        ))

    def test_failed_product_repair_still_renders_without_false_model_rank(self):
        tables, flags = self._tables_and_flags()
        invalid = self._valid_structured_market(tables, "Australia")
        invalid["priority_plan"]["products"] = []
        responses = iter((json.dumps(invalid), json.dumps(invalid)))

        raw = orchestrator.call_validated_json(
            lambda _: next(responses), "prompt",
            lambda value: orchestrator.validate_market_json(value, tables, "Australia"),
        )
        recommendation, warnings = orchestrator.normalize_market_recommendation(
            raw, tables, "Australia", flags
        )
        pack = orchestrator.render_market_pack(tables, recommendation, flags)

        self.assertTrue(raw.get("_validation_hints"))
        self.assertTrue(warnings)
        self.assertIn("AI priority decision unavailable", pack)
        self.assertIn("| Product | — |", pack)
        self.assertIn("is not an AI recommendation", pack)

    def test_failed_semantic_repair_becomes_hint_and_still_renders(self):
        tables, flags = self._tables_and_flags()
        invalid = self._valid_structured_market(tables, "Australia")
        invalid["crm"]["weekly_plan"] = invalid["crm"]["weekly_plan"][:1]
        responses = iter((json.dumps(invalid), json.dumps(invalid)))

        raw = orchestrator.call_validated_json(
            lambda _: next(responses), "prompt",
            lambda value: orchestrator.validate_market_json(value, tables, "Australia"),
        )
        recommendation, warnings = orchestrator.normalize_market_recommendation(
            raw, tables, "Australia", flags
        )
        pack = orchestrator.render_market_pack(tables, recommendation, flags)

        self.assertTrue(raw.get("_validation_hints"))
        self.assertTrue(warnings)
        self.assertIn("Campaign cadence", pack)
        self.assertEqual(len(recommendation["crm"]["weekly_plan"]),
                         len(orchestrator.campaign_weeks(tables)))

    def test_invalid_json_after_repair_reports_model_failure_without_fallback(self):
        tables, _ = self._tables_and_flags()
        responses = iter(("not json", "still not json"))

        with self.assertRaisesRegex(orchestrator.StructuredOutputError,
                                    "Model returned invalid JSON twice"):
            orchestrator.call_validated_json(
                lambda _: next(responses), "prompt",
                lambda value: orchestrator.validate_market_json(value, tables, "Australia"),
            )

    def test_model_call_failure_is_reported_without_fallback(self):
        def fail(_):
            raise RuntimeError("provider unavailable")

        with self.assertRaisesRegex(RuntimeError, "provider unavailable"):
            orchestrator.call_validated_json(fail, "prompt", lambda _: [])

    def test_openai_compatible_stream_hides_reasoning_and_returns_final_content(self):
        captured = {}

        class FakeOpenAI:
            def __init__(self, **kwargs):
                captured["client"] = kwargs

            def __enter__(self):
                return self

            def __exit__(self, *_):
                captured["client"]["http_client"].close()

            @property
            def chat(self):
                def create(**kwargs):
                    captured["request"] = kwargs
                    return iter([
                        SimpleNamespace(choices=[SimpleNamespace(delta=SimpleNamespace(
                            reasoning_content="private reasoning", content=None))]),
                        SimpleNamespace(choices=[SimpleNamespace(delta=SimpleNamespace(
                            reasoning_content=None, content='{"ok":'))]),
                        SimpleNamespace(choices=[SimpleNamespace(delta=SimpleNamespace(
                            reasoning_content=None, content="true}"))]),
                    ])
                return SimpleNamespace(completions=SimpleNamespace(create=create))

        statuses = []
        with patch("core.llm.OpenAI", FakeOpenAI):
            result = llm.generate(
                "DeepSeek", "https://example.test", "key", "model",
                "system", "user", thinking=True, status_callback=statuses.append,
            )

        self.assertEqual(result, '{"ok":true}')
        self.assertNotIn("private reasoning", result)
        self.assertTrue(captured["request"]["stream"])
        self.assertEqual(captured["request"]["reasoning_effort"], "high")
        self.assertEqual(captured["client"]["timeout"].read, 600.0)
        self.assertTrue(any("thinking stream active" in status for status in statuses))
        self.assertTrue(any("response received" in status for status in statuses))

    def test_stream_usage_is_captured_from_final_empty_choices_chunk(self):
        captured = {}

        class FakeOpenAI:
            def __init__(self, **kwargs):
                self.http_client = kwargs["http_client"]

            def __enter__(self): return self

            def __exit__(self, *_): self.http_client.close()

            @property
            def chat(self):
                def create(**kwargs):
                    captured["request"] = kwargs
                    return iter([
                        SimpleNamespace(choices=[SimpleNamespace(delta=SimpleNamespace(
                            reasoning_content="hidden", content=None))], usage=None),
                        SimpleNamespace(choices=[SimpleNamespace(delta=SimpleNamespace(
                            reasoning_content=None, content='{"ok":true}'))], usage=None),
                        SimpleNamespace(choices=[], usage=SimpleNamespace(
                            prompt_tokens=120, completion_tokens=30, total_tokens=150,
                            completion_tokens_details=SimpleNamespace(reasoning_tokens=18),
                        )),
                    ])
                return SimpleNamespace(completions=SimpleNamespace(create=create))

        usage_rows = []
        with patch("core.llm.OpenAI", FakeOpenAI):
            result = llm.generate(
                "Qwen (DashScope)", "https://example.test", "key", "model",
                "system", "user", thinking=True, usage_callback=usage_rows.append,
            )

        self.assertEqual(result, '{"ok":true}')
        self.assertEqual(captured["request"]["stream_options"], {"include_usage": True})
        self.assertEqual(usage_rows[0]["total_tokens"], 150)
        self.assertEqual(usage_rows[0]["reasoning_tokens"], 18)
        self.assertGreater(usage_rows[0]["reasoning_chars"], 0)

    def test_semantic_repair_accepts_partial_patch_and_preserves_valid_fields(self):
        tables, _ = self._tables_and_flags()
        invalid = self._valid_structured_market(tables, "Australia")
        invalid["market_summary"] = "Keep this distinctive valid market summary unchanged"
        invalid["crm"]["weekly_plan"] = invalid["crm"]["weekly_plan"][:1]
        valid = self._valid_structured_market(tables, "Australia")
        responses = iter((json.dumps(invalid), json.dumps({"crm": valid["crm"]})))
        calls = []

        repaired = orchestrator.call_validated_json(
            lambda prompt: calls.append(prompt) or next(responses),
            "large original prompt that must not be repeated",
            lambda value: orchestrator.validate_market_json(value, tables, "Australia"),
            repair_context="compact legal choices",
        )

        self.assertEqual(repaired["market_summary"], invalid["market_summary"])
        self.assertEqual(repaired["crm"], valid["crm"])
        self.assertNotIn("large original prompt", calls[1])
        self.assertIn("compact legal choices", calls[1])
        self.assertIn("these repair units: ['crm']", calls[1])

    def test_market_prompt_uses_catalog_once_instead_of_repeating_source_tables(self):
        tables, flags = self._tables_and_flags()
        package, _ = context.assemble_market_prompt(tables, "Australia", flags)
        catalog = orchestrator.market_evidence_catalog(tables, "Australia")
        prompt = prompts.MARKET_RECOMMENDATION_JSON.format(
            market="Australia", weeks=orchestrator.campaign_weeks(tables),
            evidence_catalog=json.dumps(catalog, ensure_ascii=False, separators=(",", ":")),
            context=package,
        )

        self.assertEqual(prompt.count("Booking Lead Time ="), 1)
        self.assertNotIn("Market demand (table B", prompt)
        self.assertIn('"legal_choices"', prompt)

    def test_non_thinking_path_also_uses_streaming(self):
        captured = {}

        class FakeOpenAI:
            def __init__(self, **kwargs):
                self.http_client = kwargs["http_client"]

            def __enter__(self): return self

            def __exit__(self, *_): self.http_client.close()

            @property
            def chat(self):
                def create(**kwargs):
                    captured.update(kwargs)
                    return iter([SimpleNamespace(choices=[SimpleNamespace(
                        delta=SimpleNamespace(reasoning_content=None, content="done"))])])
                return SimpleNamespace(completions=SimpleNamespace(create=create))

        with patch("core.llm.OpenAI", FakeOpenAI):
            result = llm.generate(
                "DeepSeek", "https://example.test", "key", "model",
                "system", "user", thinking=False,
            )

        self.assertEqual(result, "done")
        self.assertTrue(captured["stream"])
        self.assertEqual(captured["extra_body"], {"thinking": {"type": "disabled"}})
        self.assertNotIn("reasoning_effort", captured)

    def test_crm_normalization_auto_adds_channel_and_city_evidence(self):
        tables, flags = self._tables_and_flags()
        raw = self._valid_structured_market(tables, "Australia")
        raw["crm"]["weekly_plan"][0]["action"] = "Activate Tokyo through the selected channel"
        raw["crm"]["weekly_plan"][0]["evidence_refs"] = []

        recommendation, _ = orchestrator.normalize_market_recommendation(
            raw, tables, "Australia", flags
        )
        displays = [item["display"] for item in
                    recommendation["crm"]["weekly_plan"][0]["evidence_basis"]]

        self.assertTrue(any("Tokyo role =" in display for display in displays))
        self.assertTrue(any("eligibility = eligible" in display for display in displays))

    def test_evidence_catalog_exposes_notes_and_performance_attributes(self):
        tables, _ = parser.load_workbook_with_report(ROOT / "data" / "go_china.xlsx")
        catalog = orchestrator.market_evidence_catalog(tables, "Thailand")

        channel = catalog["E.CHANNEL.CONTENT_CHANNEL"]["display"]
        module = catalog["F.MODULE.HERO_KV"]["display"]
        asset = catalog["G.ASSET.HOMEPAGE_BANNER"]["display"]
        city = catalog["DERIVED.CITY_ROLE.GUANGZHOU"]["display"]
        city_readiness = catalog["D.CITY_READINESS.GUANGZHOU"]["display"]
        constraint = catalog["H.CONSTRAINT.CRM_TEAM"]["display"]

        self.assertIn("CTR Index", channel)
        self.assertIn("Notes", channel)
        self.assertIn("Design Complexity", module)
        self.assertIn("Localization Level", asset)
        self.assertIn("campaign risk", city)
        self.assertIn("Destination Role", city_readiness)
        self.assertIn("Attractions / Tours Readiness", city_readiness)
        self.assertIn("Impact on AI Workflow", constraint)

    def test_multiple_stakeholder_constraints_are_preserved_and_routed(self):
        tables, flags = self._tables_and_flags()
        raw = self._valid_structured_market(tables, "Australia")
        contract = next(item for item in plan.recommendation_contract(tables)["markets"]
                        if item["market"] == "Australia")
        asset = contract["assets"]["eligible"][0]["_name"]
        raw["cities"][0]["assets"] = [{
            "id": asset,
            "justification": "Supports the selected market page and localization job",
        }]
        raw["priority_plan"]["assets"] = [{
            "id": asset,
            "decision": "Produce this asset for the selected market experience",
            "reason": "Its source usage matches the selected page and localization job",
            "evidence_refs": [f"G.ASSET.{orchestrator._slug(asset)}"],
        }]
        recommendation, _ = orchestrator.normalize_market_recommendation(
            raw, tables, "Australia", flags
        )

        names = {item["stakeholder"] for item in recommendation["stakeholder_constraints"]}
        self.assertIn("Regional Teams", names)
        self.assertIn("Sourcing Team", names)
        self.assertIn("CRM Team", names)
        pack = orchestrator.render_market_pack(tables, recommendation, flags)
        self.assertIn("Applicable stakeholder constraints", pack)
        self.assertIn("Push frequency is limited", pack)
        asset_section = pack.split("## 5. Asset requests", 1)[1].split(
            "## 6. Localization checklist", 1
        )[0]
        self.assertLess(asset_section.index("Applicable stakeholder constraints"),
                        asset_section.index("| Asset | Priority | Why selected [AI REC] |"))
        self.assertIn(
            "| Asset | Priority | Why selected [AI REC] |\n|---|---|---|\n| " + asset + " |",
            asset_section,
        )

    def test_confirmation_ids_render_with_semantic_labels(self):
        tables, flags = self._tables_and_flags()
        recommendation, _ = orchestrator.normalize_market_recommendation(
            {"cities": []}, tables, "Australia", flags
        )
        pack = orchestrator.render_market_pack(tables, recommendation, flags)

        first_relevant = orchestrator._relevant_flags(flags, recommendation)[0]
        self.assertIn(f"{first_relevant['_id']} —", pack)


if __name__ == "__main__":
    unittest.main()
